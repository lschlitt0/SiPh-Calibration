"""_write_xlsx extracted from the calibration experiment runner."""

from __future__ import annotations

from pathlib import Path
from typing import Any
from typing import Dict
from typing import Sequence
from typing import Tuple

from ._xlsx_cell import _xlsx_cell
from .log import log


def _write_xlsx(path: Path, sheets: Dict[str, Tuple[Sequence[str], Sequence[Dict[str, Any]]]]) -> None:
    """Export named result sheets using openpyxl or, failing import, XlsxWriter.

    If neither package imports, log the omission and return without a workbook.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        from openpyxl import Workbook
    except Exception:
        Workbook = None
    if Workbook is not None:
        wb = Workbook(write_only=True)
        for name, (fieldnames, rows) in sheets.items():
            sheet_name = str(name)[:31]
            ws = wb.create_sheet(title=sheet_name)
            ws.append(list(fieldnames))
            for row in rows:
                ws.append([_xlsx_cell(row.get(k)) for k in fieldnames])
        wb.save(path)
        return
    try:
        import xlsxwriter
    except Exception as exc:
        log(f"[xlsx] missing openpyxl/xlsxwriter ({exc}); skipping workbook", force=True)
        return
    wb = xlsxwriter.Workbook(str(path))
    try:
        for name, (fieldnames, rows) in sheets.items():
            sheet_name = str(name)[:31]
            ws = wb.add_worksheet(sheet_name)
            ws.write_row(0, 0, list(fieldnames))
            r = 1
            for row in rows:
                ws.write_row(r, 0, [_xlsx_cell(row.get(k)) for k in fieldnames])
                r += 1
    finally:
        wb.close()
